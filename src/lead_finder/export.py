"""Write lead lists to CSV / Excel and compute coverage stats."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .clean import site_domain
from .models import Business

COLUMNS = [
    "name", "website", "email", "email_source", "email_matches_site", "email_mx_ok", "phone", "all_emails", "all_phones",
    "facebook", "instagram", "linkedin", "x", "youtube", "address", "city", "state", "postcode",
    "lat", "lon", "crawl_status", "osm_url",
]


def to_row(biz: Business) -> dict:
    return {
        "name": biz.name,
        "website": biz.website,
        "email": biz.best_email,
        "email_source": biz.email_source,
        "email_matches_site": _matches_site(biz.best_email, biz.website) if biz.best_email else "",
        "email_mx_ok": "" if biz.email_mx_ok is None else biz.email_mx_ok,
        "phone": biz.best_phone,
        "all_emails": "; ".join(biz.all_emails),
        "all_phones": "; ".join(biz.all_phones),
        **{network: biz.socials.get(network, "") for network in ("facebook", "instagram", "linkedin", "x", "youtube")},
        "address": biz.address,
        "city": biz.city,
        "state": biz.state,
        "postcode": biz.postcode,
        "lat": biz.lat,
        "lon": biz.lon,
        "crawl_status": biz.crawl_status,
        "osm_url": biz.osm_url,
    }


def _matches_site(email: str, website: str) -> bool:
    domain = site_domain(website)
    email_domain = email.partition("@")[2]
    return bool(domain) and (email_domain == domain or email_domain.endswith("." + domain))


def coverage(businesses: list[Business]) -> dict[str, int]:
    """Counts used for the run summary (OSM-only vs. after website enrichment)."""
    return {
        "businesses": len(businesses),
        "with_website": sum(1 for b in businesses if b.website),
        "crawled_ok": sum(1 for b in businesses if b.crawl_status == "ok"),
        "email_from_osm": sum(1 for b in businesses if b.osm_emails),
        "email_after_enrichment": sum(1 for b in businesses if b.best_email),
        "phone_from_osm": sum(1 for b in businesses if b.osm_phones),
        "phone_after_enrichment": sum(1 for b in businesses if b.best_phone),
        "with_any_social": sum(1 for b in businesses if b.socials),
    }


def write_csv(businesses: list[Business], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=COLUMNS)
        writer.writeheader()
        writer.writerows(to_row(b) for b in businesses)


def write_xlsx(businesses: list[Business], path: Path, stats: dict[str, int]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append(COLUMNS)
    for biz in businesses:
        row = to_row(biz)
        ws.append([row[c] for c in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for idx, col in enumerate(COLUMNS, start=1):
        width = max(len(col), *(len(str(c.value or "")) for c in ws[get_column_letter(idx)][1:200]))
        ws.column_dimensions[get_column_letter(idx)].width = min(width + 2, 45)

    summary = wb.create_sheet("Summary")
    summary.append(["metric", "count"])
    for key, value in stats.items():
        summary.append([key, value])
    summary["A1"].font = summary["B1"].font = Font(bold=True)
    summary.column_dimensions["A"].width = 28
    wb.save(path)
