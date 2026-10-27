import csv

from openpyxl import load_workbook

from lead_finder.clean import finalize
from lead_finder.export import COLUMNS, coverage, to_row, write_csv, write_xlsx
from lead_finder.models import Business


def sample() -> list[Business]:
    a = Business(osm_type="node", osm_id=1, name="Acme HVAC", website="https://acmehvac.com",
                 osm_phones=["305-555-0142"], site_emails=["info@acmehvac.com"],
                 socials={"facebook": "https://facebook.com/acme"}, crawl_status="ok")
    b = Business(osm_type="way", osm_id=2, name="Bay Air", site_emails=["bayair@gmail.com"],
                 website="https://bayair.com", crawl_status="ok")
    c = Business(osm_type="node", osm_id=3, name="No Contact Co", crawl_status="no_website")
    return [finalize(x) for x in (a, b, c)]


def test_to_row_flags_email_domain_match():
    a, b, c = (to_row(x) for x in sample())
    assert a["email_matches_site"] is True and a["facebook"] == "https://facebook.com/acme"
    assert b["email_matches_site"] is False
    assert c["email_matches_site"] == "" and c["email"] == ""


def test_coverage_counts():
    stats = coverage(sample())
    assert stats["businesses"] == 3
    assert stats["with_website"] == 2
    assert stats["email_from_osm"] == 0
    assert stats["email_after_enrichment"] == 2
    assert stats["phone_after_enrichment"] == 1


def test_write_csv_and_xlsx(tmp_path):
    items = sample()
    write_csv(items, tmp_path / "leads.csv")
    write_xlsx(items, tmp_path / "leads.xlsx", coverage(items))

    with (tmp_path / "leads.csv").open(encoding="utf-8-sig") as f:
        rows = list(csv.DictReader(f))
    assert list(rows[0].keys()) == COLUMNS and len(rows) == 3

    wb = load_workbook(tmp_path / "leads.xlsx")
    assert wb.sheetnames == ["Leads", "Summary"]
    assert wb["Leads"].max_row == 4
    assert wb["Summary"]["A2"].value == "businesses"
